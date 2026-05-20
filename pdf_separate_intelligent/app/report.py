import json
import os
import csv
import io
from typing import Dict, Optional


def load_report_results(output_dir: str) -> Optional[Dict]:
    report_path = os.path.join(output_dir, "_report_results.json")
    if not os.path.exists(report_path):
        return None
    with open(report_path, "r", encoding="utf-8") as f:
        return json.load(f)


def generate_csv_report(report_data: Dict) -> str:
    output = io.StringIO()
    writer = csv.writer(output)

    employees = report_data.get("employees", {})
    total_expected = report_data.get("total_expected", 0)
    total_processed = report_data.get("total_processed", 0)
    total_failed = report_data.get("total_failed", 0)

    # Summary section
    writer.writerow(["=== RAPPORT D'IMPORTATION & RÉCONCILIATION ==="])
    writer.writerow([])
    writer.writerow(["Indicateur", "Valeur"])
    writer.writerow(["Total attendu (pages)", total_expected])
    writer.writerow(["Total traités avec succès", total_processed])
    writer.writerow(["Total en échec", total_failed])
    writer.writerow(["Taux de succès", f"{total_processed / total_expected * 100:.1f}%" if total_expected > 0 else "N/A"])
    writer.writerow([])
    writer.writerow([])

    # Header
    writer.writerow(["Matricule", "Nom", "Prénom", "Fichier", "Statut", "Raison d'échec"])

    # Data rows grouped by employee
    for mat in sorted(employees.keys(), key=lambda x: (x.startswith("UNKNOWN"), x)):
        emp = employees[mat]
        first_row = True
        for f in emp["files"]:
            writer.writerow([
                emp["matricule"] if first_row else "",
                emp["nom"] if first_row else "",
                emp["prenom"] if first_row else "",
                f["filename"],
                f["status"],
                f["error"] or ""
            ])
            first_row = False
        # Employee summary
        writer.writerow([
            "", "", "",
            f"Total pour {emp['matricule']}: {emp['total']} fichier(s), {emp['success']} succès, {emp['failed']} échec(s)",
            "", ""
        ])
        writer.writerow([])

    return output.getvalue()


def generate_xlsx_report(report_data: Dict) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl is required for XLSX generation. Install it with: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport Importation"

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="166534", end_color="166534", fill_type="solid")
    title_font = Font(bold=True, size=14, color="166534")
    subtitle_font = Font(bold=True, size=11, color="374151")
    success_font = Font(color="166534", bold=True)
    error_font = Font(color="DC2626", bold=True)
    summary_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    error_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    employees = report_data.get("employees", {})
    total_expected = report_data.get("total_expected", 0)
    total_processed = report_data.get("total_processed", 0)
    total_failed = report_data.get("total_failed", 0)

    # Title
    ws.merge_cells('A1:F1')
    ws['A1'] = "RAPPORT D'IMPORTATION & RÉCONCILIATION"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 30

    # Summary section
    ws.merge_cells('A3:B3')
    ws['A3'] = "Résumé de la réconciliation"
    ws['A3'].font = subtitle_font

    summary_data = [
        ("Total attendu (pages dans le PDF)", total_expected),
        ("Total traités avec succès", total_processed),
        ("Total en échec", total_failed),
        ("Taux de succès", f"{total_processed / total_expected * 100:.1f}%" if total_expected > 0 else "N/A"),
    ]
    for i, (label, value) in enumerate(summary_data):
        row = 4 + i
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A4:B4')
    ws.merge_cells('A5:B5')
    ws.merge_cells('A6:B6')
    ws.merge_cells('A7:B7')

    # Data table header
    header_row = 9
    headers = ["Matricule", "Nom", "Prénom", "Fichier", "Statut", "Raison d'échec"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    # Column widths
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 55

    # Data rows grouped by employee
    current_row = header_row + 1
    for mat in sorted(employees.keys(), key=lambda x: (x.startswith("UNKNOWN"), x)):
        emp = employees[mat]
        start_row = current_row

        for f in emp["files"]:
            is_error = f["status"] == "Error"
            row_fill = error_fill if is_error else summary_fill

            ws.cell(row=current_row, column=1, value=emp["matricule"]).border = thin_border
            ws.cell(row=current_row, column=2, value=emp["nom"]).border = thin_border
            ws.cell(row=current_row, column=3, value=emp["prenom"]).border = thin_border
            ws.cell(row=current_row, column=4, value=f["filename"]).border = thin_border

            status_cell = ws.cell(row=current_row, column=5, value=f["status"])
            status_cell.font = success_font if f["status"] == "Success" else error_font
            status_cell.border = thin_border

            error_cell = ws.cell(row=current_row, column=6, value=f["error"] or "")
            error_cell.border = thin_border

            for col in range(1, 7):
                ws.cell(row=current_row, column=col).fill = row_fill

            current_row += 1

        # Merge employee info cells across their rows
        if current_row - start_row > 1:
            ws.merge_cells(start_row=start_row, start_column=1, end_row=current_row - 1, end_column=1)
            ws.merge_cells(start_row=start_row, start_column=2, end_row=current_row - 1, end_column=2)
            ws.merge_cells(start_row=start_row, start_column=3, end_row=current_row - 1, end_column=3)

        # Employee summary line
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        summary_cell = ws.cell(row=current_row, column=1,
                               value=f"Total pour {emp['matricule']}: {emp['total']} fichier(s), {emp['success']} succès, {emp['failed']} échec(s)")
        summary_cell.font = Font(bold=True, italic=True, color="374151")
        summary_cell.border = thin_border
        current_row += 2

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
